"""
DOCX Extraction Engine — Section 3 & Section 9 of Master Implementation Specification.

Uses python-docx to extract headings, paragraphs, and tables from Word documents.
"""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Dict, List, Optional

import docx

from core.domain.documents import (
    CanonicalDocument,
    DocumentElement,
    DocumentType,
    ElementType,
    Page,
)
from core.extraction.base import BaseExtractor
from core.ingestion.discovery import DiscoveredFile
from core.provenance.tracker import generate_document_id, generate_element_id


class DOCXExtractor(BaseExtractor):
    """Extracts structured content from Microsoft Word documents."""

    def extract(self, file_path: Path, discovered: Optional[DiscoveredFile] = None) -> CanonicalDocument:
        path = Path(file_path)
        with open(path, "rb") as f:
            raw_bytes = f.read()

        file_size = len(raw_bytes)
        source_hash = hashlib.sha256(raw_bytes).hexdigest()
        doc_id = generate_document_id(str(path), raw_bytes)

        doc = docx.Document(str(path))
        elements: List[DocumentElement] = []
        tables_data: List[Dict[str, Any]] = []
        markdown_sections: List[str] = []

        counter = 0

        # Extract paragraphs & headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            counter += 1
            style_name = p.style.name.lower() if p.style else ""
            is_heading = "heading" in style_name or "title" in style_name
            el_type = ElementType.HEADING if is_heading else ElementType.PARAGRAPH
            el_id = generate_element_id(doc_id, 1, counter)

            elements.append(
                DocumentElement(
                    element_id=el_id,
                    document_id=doc_id,
                    type=el_type,
                    page_number=1,
                    text=text,
                    reading_order=counter,
                    confidence=1.0,
                    metadata={"style": style_name},
                )
            )

            if is_heading:
                markdown_sections.append(f"## {text}\n")
            else:
                markdown_sections.append(f"{text}\n")

        # Extract tables
        for tbl_idx, table in enumerate(doc.tables, start=1):
            table_rows: List[List[str]] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_cells)

            if table_rows:
                counter += 1
                el_id = generate_element_id(doc_id, 1, counter)
                headers = table_rows[0]
                body_rows = table_rows[1:] if len(table_rows) > 1 else []

                tables_data.append({
                    "table_index": tbl_idx,
                    "headers": headers,
                    "rows": body_rows,
                })

                elements.append(
                    DocumentElement(
                        element_id=el_id,
                        document_id=doc_id,
                        type=ElementType.TABLE,
                        page_number=1,
                        text=f"Table with {len(table_rows)} rows and {len(headers)} columns",
                        reading_order=counter,
                        confidence=1.0,
                        metadata={"table_index": tbl_idx, "headers": headers},
                    )
                )

                header_line = "| " + " | ".join(headers) + " |"
                sep_line = "| " + " | ".join(["---"] * len(headers)) + " |"
                markdown_sections.append(header_line)
                markdown_sections.append(sep_line)
                for r in body_rows[:20]:
                    markdown_sections.append("| " + " | ".join(r) + " |")
                markdown_sections.append("\n")

        # Inherit temporal metadata
        fy = discovered.temporal.financial_year if discovered else None
        period = discovered.temporal.reporting_period if discovered else None
        dates = discovered.temporal.extracted_dates if discovered else []

        page = Page(page_number=1, elements=elements)

        return CanonicalDocument(
            document_id=doc_id,
            source_reference=str(path.resolve()),
            source_hash=source_hash,
            document_type=DocumentType.DOCX,
            reporting_year=fy,
            reporting_period=period,
            dates=dates,
            file_size_bytes=file_size,
            pages=[page],
            tables=tables_data,
            markdown_content="\n".join(markdown_sections),
            metadata={"total_paragraphs": len(doc.paragraphs), "total_tables": len(doc.tables)},
        )
