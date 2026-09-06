"""
XLSX Extraction Engine — Section 9 of Master Implementation Specification.

Uses openpyxl to extract workbooks, sheets, and individual cell coordinates.
Generates exact cell-level provenance (e.g. Workbook: coal_stats.xlsx, Sheet: March, Cell: G27, Value: 1245.70).
"""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from core.domain.documents import (
    CanonicalDocument,
    DocumentElement,
    DocumentType,
    ElementType,
    Page,
)
from core.domain.evidence import SpreadsheetCoordinate
from core.extraction.base import BaseExtractor
from core.ingestion.discovery import DiscoveredFile
from core.provenance.tracker import generate_document_id, generate_element_id


class XLSXExtractor(BaseExtractor):
    """
    Extracts structured spreadsheets into tables and cell-level elements.
    Preserves exact workbook, sheet, and cell coordinate provenance.
    """

    def extract(self, file_path: Path, discovered: Optional[DiscoveredFile] = None) -> CanonicalDocument:
        path = Path(file_path)
        with open(path, "rb") as f:
            raw_bytes = f.read()

        file_size = len(raw_bytes)
        source_hash = hashlib.sha256(raw_bytes).hexdigest()
        doc_id = generate_document_id(str(path), raw_bytes)
        workbook_name = path.name

        wb = openpyxl.load_workbook(filename=str(path), data_only=True, read_only=True)
        sheet_names = wb.sheetnames

        pages: List[Page] = []
        all_tables: List[Dict[str, Any]] = []
        markdown_sections: List[str] = [f"# Workbook: {workbook_name}\n"]
        element_counter = 0

        for sheet_idx, sheet_name in enumerate(sheet_names):
            sheet = wb[sheet_name]
            sheet_elements: List[DocumentElement] = []
            rows_data: List[List[Any]] = []

            markdown_sections.append(f"## Sheet: {sheet_name}\n")

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                row_values = list(row)
                # Skip entirely empty rows
                if not any(v is not None for v in row_values):
                    continue

                rows_data.append(row_values)

                for col_idx, val in enumerate(row_values, start=1):
                    if val is None or str(val).strip() == "":
                        continue

                    col_letter = openpyxl.utils.get_column_letter(col_idx)
                    cell_coord_str = f"{col_letter}{row_idx}"
                    element_counter += 1
                    el_id = generate_element_id(doc_id, sheet_idx + 1, element_counter)

                    formatted_str = str(val).strip()
                    sheet_coord = SpreadsheetCoordinate(
                        workbook_name=workbook_name,
                        sheet_name=sheet_name,
                        cell=cell_coord_str,
                        row=row_idx,
                        column=col_idx,
                        raw_value=val,
                        formatted_value=formatted_str,
                    )

                    sheet_elements.append(
                        DocumentElement(
                            element_id=el_id,
                            document_id=doc_id,
                            type=ElementType.SPREADSHEET_CELL,
                            page_number=sheet_idx + 1,
                            text=formatted_str,
                            confidence=1.0,
                            metadata={
                                "spreadsheet_coord": sheet_coord.model_dump(),
                                "raw_value": val,
                            },
                        )
                    )

            # Record table structure for this sheet
            if rows_data:
                headers = [str(c or "") for c in rows_data[0]]
                body_rows = rows_data[1:] if len(rows_data) > 1 else []
                all_tables.append({
                    "sheet_name": sheet_name,
                    "headers": headers,
                    "row_count": len(rows_data),
                    "col_count": len(headers),
                    "rows": [[str(c if c is not None else "") for c in r] for r in body_rows[:100]],
                })

                # Format sheet as markdown table (first 25 rows for LLM context)
                if headers:
                    header_line = "| " + " | ".join(headers) + " |"
                    sep_line = "| " + " | ".join(["---"] * len(headers)) + " |"
                    markdown_sections.append(header_line)
                    markdown_sections.append(sep_line)
                    for r in body_rows[:25]:
                        row_line = "| " + " | ".join(str(c if c is not None else "") for c in r) + " |"
                        markdown_sections.append(row_line)
                    markdown_sections.append("\n")

            pages.append(
                Page(
                    page_number=sheet_idx + 1,
                    elements=sheet_elements,
                    metadata={"sheet_name": sheet_name, "row_count": len(rows_data)},
                )
            )

        wb.close()

        # Inherit temporal metadata
        fy = discovered.temporal.financial_year if discovered else None
        period = discovered.temporal.reporting_period if discovered else None
        dates = discovered.temporal.extracted_dates if discovered else []

        return CanonicalDocument(
            document_id=doc_id,
            source_reference=str(path.resolve()),
            source_hash=source_hash,
            document_type=DocumentType.XLSX,
            reporting_year=fy,
            reporting_period=period,
            dates=dates,
            file_size_bytes=file_size,
            pages=pages,
            tables=all_tables,
            markdown_content="\n".join(markdown_sections),
            metadata={
                "workbook_name": workbook_name,
                "sheet_names": sheet_names,
                "total_sheets": len(sheet_names),
            },
        )
