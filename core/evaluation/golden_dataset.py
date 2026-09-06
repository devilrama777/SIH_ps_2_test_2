"""
Golden Dataset Builder — Section 30 of Master Plan.

Generates the structured regression dataset under testdata/reference_report/
representing realistic CIL subsidiary annual report inputs (PDF/TXT, XLSX, DOCX, PNG).
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Any

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import docx
except ImportError:
    docx = None

try:
    from PIL import Image, ImageDraw
except ImportError:
    Image = None
    ImageDraw = None


class GoldenDatasetBuilder:
    """
    Creates and manages golden fixtures for full regression validation.
    """

    def __init__(self, base_dir: str = "testdata/reference_report"):
        self.base_dir = Path(base_dir)

    def build_dataset(self) -> Path:
        """Create all golden fixture files and ground truth specifications."""
        self.base_dir.mkdir(parents=True, exist_ok=True)
        (self.base_dir / "digital_pages").mkdir(parents=True, exist_ok=True)
        (self.base_dir / "financial_tables").mkdir(parents=True, exist_ok=True)
        (self.base_dir / "csr_sections").mkdir(parents=True, exist_ok=True)
        (self.base_dir / "audit_sections").mkdir(parents=True, exist_ok=True)
        (self.base_dir / "photographs").mkdir(parents=True, exist_ok=True)

        self._build_digital_pages()
        self._build_financial_tables()
        self._build_csr_sections()
        self._build_audit_sections()
        self._build_photographs()
        self._build_ground_truth()

        return self.base_dir

    def _build_digital_pages(self) -> None:
        path = self.base_dir / "digital_pages" / "CCL_Annual_Report_FY24_Highlights.txt"
        content = (
            "CENTRAL COALFIELDS LIMITED (A Subsidiary of Coal India Limited)\n"
            "ANNUAL OPERATIONAL HIGHLIGHTS — FINANCIAL YEAR 2023-24\n\n"
            "Central Coalfields Limited (CCL) achieved raw coal production of 84.50 Million Tonnes in FY 2023-24, "
            "registering a remarkable growth of 9.03% over 77.50 MT achieved in FY 2022-23.\n"
            "Total coal off-take during FY 2023-24 stood at 82.10 Million Tonnes against 75.30 MT in the preceding fiscal.\n"
            "Net turnover from operations reached Rs. 14,250.60 Crores, yielding Profit Before Tax (PBT) of Rs. 2,890.15 Crores.\n"
            "Composite Overburden Removal (OBR) scaled to 132.40 Million Cubic Metres (M.Cu.M), reflecting robust operational discipline."
        )
        path.write_text(content, encoding="utf-8")

    def _build_financial_tables(self) -> None:
        path = self.base_dir / "financial_tables" / "CCL_Production_Offtake_FY24.xlsx"
        if openpyxl:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Production_Summary"
            ws.append(["Mine_Area", "FY23_Actual", "FY24_Target", "FY24_Actual", "Growth_Percent"])
            ws.append(["North Karanpura", 28.5, 30.0, 31.2, 9.47])
            ws.append(["Barka Sayal", 14.2, 15.0, 15.1, 6.34])
            ws.append(["Rajrappa", 9.8, 10.5, 10.4, 6.12])
            ws.append(["Piparwar", 18.0, 19.0, 19.5, 8.33])
            ws.append(["Argada", 7.0, 8.0, 8.3, 18.57])
            ws.append(["Total_Production", 77.5, 82.5, 84.5, 9.03])
            wb.save(str(path))
        else:
            # Fallback CSV representation if openpyxl not available
            csv_path = self.base_dir / "financial_tables" / "CCL_Production_Offtake_FY24.csv"
            csv_path.write_text(
                "Mine_Area,FY23_Actual,FY24_Target,FY24_Actual,Growth_Percent\n"
                "North Karanpura,28.5,30.0,31.2,9.47\n"
                "Barka Sayal,14.2,15.0,15.1,6.34\n"
                "Rajrappa,9.8,10.5,10.4,6.12\n"
                "Piparwar,18.0,19.0,19.5,8.33\n"
                "Argada,7.0,8.0,8.3,18.57\n"
                "Total_Production,77.5,82.5,84.5,9.03\n",
                encoding="utf-8",
            )

    def _build_csr_sections(self) -> None:
        path = self.base_dir / "csr_sections" / "CSR_Community_Development_FY24.docx"
        if docx:
            doc = docx.Document()
            doc.add_heading("Corporate Social Responsibility & Sustainability FY 2023-24", level=1)
            p1 = doc.add_paragraph(
                "In FY 2023-24, Central Coalfields Limited deployed Rs. 124.50 Crores towards sustainable community initiatives, "
                "surpassing the statutory mandate of Rs. 118.20 Crores under Section 135 of the Companies Act."
            )
            p2 = doc.add_paragraph(
                "Flagship healthcare programs provided mobile medical vans covering 182 remote tribal villages. "
                "The educational initiative 'CCL ke Lal & Laadli' supported 120 students with free residential coaching for IIT/NEET entrance examinations."
            )
            doc.save(str(path))
        else:
            txt_path = self.base_dir / "csr_sections" / "CSR_Community_Development_FY24.txt"
            txt_path.write_text(
                "In FY 2023-24, Central Coalfields Limited deployed Rs. 124.50 Crores towards CSR initiatives, "
                "surpassing statutory mandate of Rs. 118.20 Crores.\n"
                "Flagship initiative 'CCL ke Lal & Laadli' supported 120 students with residential coaching.",
                encoding="utf-8",
            )

    def _build_audit_sections(self) -> None:
        path = self.base_dir / "audit_sections" / "CAG_Audit_Compliance_FY24.txt"
        content = (
            "COMMENTS OF THE COMPTROLLER AND AUDITOR GENERAL OF INDIA\n"
            "UNDER SECTION 143(6)(b) OF THE COMPANIES ACT, 2013 ON THE STANDALONE FINANCIAL STATEMENTS\n"
            "OF CENTRAL COALFIELDS LIMITED FOR THE YEAR ENDED 31 MARCH 2024\n\n"
            "The Comptroller and Auditor General of India has conducted supplementary audit under section 143(6)(a) "
            "of the Act of the standalone financial statements of Central Coalfields Limited for the year ended 31 March 2024.\n"
            "Based on our audit, the Comptroller and Auditor General of India has issued 'Nil Comments' on the financial statements."
        )
        path.write_text(content, encoding="utf-8")

    def _build_photographs(self) -> None:
        if Image and ImageDraw:
            # 1. Mine image (800x600)
            img1 = Image.new("RGB", (800, 600), color=(50, 75, 95))
            draw1 = ImageDraw.Draw(img1)
            draw1.rectangle([50, 50, 750, 550], outline=(180, 160, 100), width=4)
            draw1.text((70, 70), "CCL Piparwar Opencast Mine — Heavy Earth Moving Machinery", fill=(240, 240, 240))
            img1.save(str(self.base_dir / "photographs" / "open_cast_mine.png"))

            # 2. Plantation image (800x600)
            img2 = Image.new("RGB", (800, 600), color=(40, 95, 55))
            draw2 = ImageDraw.Draw(img2)
            draw2.rectangle([50, 50, 750, 550], outline=(100, 200, 120), width=4)
            draw2.text((70, 70), "Eco-Restoration & Afforestation Initiative — Rajrappa Area", fill=(240, 240, 240))
            img2.save(str(self.base_dir / "photographs" / "tree_plantation_drive.png"))

    def _build_ground_truth(self) -> None:
        ground_truth = {
            "subsidiary_name": "Central Coalfields Limited",
            "reporting_period": "FY 2023-24",
            "key_metrics": {
                "raw_coal_production_fy24_mt": 84.50,
                "raw_coal_production_fy23_mt": 77.50,
                "production_growth_percent": 9.03,
                "offtake_fy24_mt": 82.10,
                "offtake_fy23_mt": 75.30,
                "net_turnover_cr": 14250.60,
                "profit_before_tax_cr": 2890.15,
                "obr_m_cum": 132.40,
                "csr_expenditure_cr": 124.50,
                "csr_statutory_mandate_cr": 118.20,
                "cag_comments": "Nil Comments",
            },
            "source_files": [
                "CCL_Annual_Report_FY24_Highlights.txt",
                "CCL_Production_Offtake_FY24.xlsx",
                "CSR_Community_Development_FY24.docx",
                "CAG_Audit_Compliance_FY24.txt",
            ],
        }
        gt_path = self.base_dir / "ground_truth.json"
        gt_path.write_text(json.dumps(ground_truth, indent=2), encoding="utf-8")
