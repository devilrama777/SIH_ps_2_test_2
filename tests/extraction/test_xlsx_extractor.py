"""
Tests for XLSXExtractor (Section 9 spreadsheet coordinate provenance).
"""
from pathlib import Path
import openpyxl
from core.domain.documents import DocumentType, ElementType
from core.extraction.xlsx_extractor import XLSXExtractor


def test_xlsx_extraction_coordinate_provenance(tmp_path: Path):
    """Verify spreadsheet extraction with exact cell coordinates like G27."""
    wb_path = tmp_path / "coal_stats.xlsx"

    wb = openpyxl.Workbook()
    # Sheet 1: March
    ws_march = wb.active
    ws_march.title = "March"
    ws_march["A1"] = "Colliery"
    ws_march["B1"] = "Production_MT"
    ws_march["A2"] = "North Karanpura"
    ws_march["B2"] = 750.50
    ws_march["G27"] = 1245.70  # Explicitly matches Section 9 example

    # Sheet 2: Summary
    ws_summary = wb.create_sheet(title="Summary")
    ws_summary["A1"] = "Total Subsidiaries"
    ws_summary["B1"] = 8

    wb.save(str(wb_path))
    wb.close()

    extractor = XLSXExtractor()
    canonical = extractor.extract(wb_path)

    assert canonical.document_type == DocumentType.XLSX
    assert len(canonical.pages) == 2
    assert len(canonical.tables) == 2

    # Find the G27 cell element
    march_page = canonical.pages[0]
    g27_elements = [
        e for e in march_page.elements
        if e.metadata.get("spreadsheet_coord", {}).get("cell") == "G27"
    ]
    assert len(g27_elements) == 1
    g27_el = g27_elements[0]

    assert g27_el.type == ElementType.SPREADSHEET_CELL
    assert g27_el.text == "1245.7"
    coord = g27_el.metadata["spreadsheet_coord"]
    assert coord["sheet_name"] == "March"
    assert coord["row"] == 27
    assert coord["column"] == 7
    assert coord["raw_value"] == 1245.70

    # Verify tables structure
    table_march = canonical.tables[0]
    assert table_march["sheet_name"] == "March"
    assert "Colliery" in table_march["headers"]
