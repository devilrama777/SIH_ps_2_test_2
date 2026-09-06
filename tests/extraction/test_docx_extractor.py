"""
Tests for DOCXExtractor.
"""
from pathlib import Path
import docx
from core.domain.documents import DocumentType, ElementType
from core.extraction.docx_extractor import DOCXExtractor


def test_docx_extraction(tmp_path: Path):
    """Verify DOCX heading, paragraph, and table extraction."""
    doc_path = tmp_path / "board_minutes.docx"

    doc = docx.Document()
    doc.add_heading("Board of Directors Review", level=1)
    doc.add_paragraph("The 340th meeting discussed safety audits in underground mines.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Safety Metric"
    table.cell(0, 1).text = "FY25 Score"
    table.cell(1, 0).text = "Fatality Rate"
    table.cell(1, 1).text = "0.02"

    doc.save(str(doc_path))

    extractor = DOCXExtractor()
    canonical = extractor.extract(doc_path)

    assert canonical.document_type == DocumentType.DOCX
    assert len(canonical.pages) == 1
    assert len(canonical.tables) == 1

    elements = canonical.pages[0].elements
    headings = [e for e in elements if e.type == ElementType.HEADING]
    paragraphs = [e for e in elements if e.type == ElementType.PARAGRAPH]
    tables = [e for e in elements if e.type == ElementType.TABLE]

    assert len(headings) >= 1
    assert "Board of Directors" in headings[0].text
    assert len(paragraphs) >= 1
    assert "safety audits" in paragraphs[0].text
    assert len(tables) >= 1
